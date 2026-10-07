from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from io import BytesIO
import zipfile
import re

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "revision_output" / "response" / "Response to Review.docx"
REFERENCE = Path(r"C:\Users\Cem\Downloads\PAJES-Access-Response-to-Reviewers.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)


def plain_text(value):
    replacements = {
        "We also reconciled the 74 non-empty archived records with the 72-configuration item-level set: record 0 is a 49-case partial suite and record 18 is an earlier auxiliary GPT-OSS/GPT-OSS Serial result whose systematically named counterpart is record 39.":
            "We also reconciled the 74 non-empty archived records with the 72-configuration item-level set. Record 0 is a 49-case partial suite. Record 18 is an earlier auxiliary GPT-OSS/GPT-OSS Serial result. Its systematically named counterpart is record 39.",
        "We agree that the archive is not a fully crossed strategy–optimizer experiment. We identified one limited same-source subset: the 169-case LLaMa3.2-3B suite evaluated with Gemini-2.5 Flash and Pro under Serial, Bulk, and Parallel. We report all six archived outcomes descriptively, while withholding a causal strategy ranking because they are single records and provider/run settings are incomplete.":
            "We agree that the archive does not test every strategy and optimizer combination. We identified one limited same-source subset. It is the 169-case LLaMa3.2-3B suite evaluated with Gemini-2.5 Flash and Pro under Serial, Bulk, and Parallel. We report all six archived outcomes as descriptions. We do not give a causal strategy ranking because the archive has one record per configuration and lacks complete provider and run settings.",
        "We report item-level precision, recall, F1, balanced accuracy, confusion counts, Cohen's kappa, false positives, and false negatives for every included configuration, and explicitly explain why the item-level set contains 72 of the 74 non-empty archived records.":
            "We report item-level precision, recall, F1, balanced accuracy, confusion counts, Cohen's kappa, false positives, and false negatives for every included configuration. We also explain why the item-level set contains 72 of the 74 non-empty archived records.",
        "The files do not document independence, blinding, adjudication, or whether labels were mutually visible, so independence is not claimed.":
            "The files do not document independence, blinding, or adjudication. They also do not show whether labels were mutually visible. We therefore do not claim independence.",
        "The archive does not establish prospective threshold selection, so we removed 'pre-specified' and do not claim that this single baseline is a comprehensive state-of-the-art comparison.":
            "The archive does not show that the thresholds were selected in advance. We therefore removed 'pre-specified'. We also do not present this single baseline as a complete state-of-the-art comparison.",
        "The figures remain legible at the journal page size, while large matrices are retained in compact form because the source image assets did not support lossless label regeneration.":
            "The figures remain legible at the journal page size. The large matrices remain compact because the source image files do not support lossless label regeneration.",
        "The original XML, requirements, and source files are not included in the supplied revision package, so exact structural-element and requirement counts were removed rather than repeated without traceable source artefacts.":
            "The supplied revision package does not include the original XML, requirements, or source files. We therefore removed exact structural-element and requirement counts that lacked traceable source artefacts.",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    value = value.replace("—", ". ")
    value = re.sub(r";\s+([a-z])", lambda match: ". " + match.group(1).upper(), value)
    value = value.replace("; ", ". ").replace(";", ".")
    return re.sub(r"\s+", " ", value).strip()


def item(comment, response, action, location, excerpt, status="Fully addressed"):
    if status.startswith("Partially addressed"):
        status = "Partially addressed"
    elif status == "Addressed":
        status = "Fully addressed"
    return dict(
        comment=plain_text(comment),
        response=plain_text(response),
        action=plain_text(action),
        location=plain_text(location),
        excerpt=plain_text(excerpt),
        status=plain_text(status),
    )


COMMON_EXECUTION = (
    "We agree that semantic redundancy is not a substitute for execution evidence. The supplied archive contains no "
    "execution outcomes, coverage traces, mutation results, or fault data, so we did not manufacture these results. "
    "We narrowed the claims and made execution-based validation an explicit downstream study."
)
COMMON_EXECUTION_EXCERPT = (
    "The generated cases were not executed. The study therefore does not establish execution success, code or branch "
    "coverage, mutation adequacy, fault detection, behavioural adequacy, or overall test-suite effectiveness."
)
COMMON_STOCHASTIC = (
    "We agree. No repeated-run records or complete inference settings were present in the supplied archive. We replaced "
    "the former repetition claim with an archive-bounded statement and explicitly identify the missing controls."
)
COMMON_STOCHASTIC_EXCERPT = (
    "Comprehensive repeated-run experiments are not available. The results therefore describe single archived outputs "
    "rather than robustness distributions; no repeated-run means, variances, confidence intervals, or significance tests are claimed."
)
COMMON_ORACLE = (
    "We re-audited the three Human Oracle directories. Three separately completed evaluator files cover the same 1,074 items. "
    "We now use a three-rater majority consensus, report Fleiss' kappa and all pairwise Cohen kappas, and provide item-level decisions and confusion metrics. "
    "The files do not document independence, blinding, adjudication, or whether labels were mutually visible, so independence is not claimed."
)
COMMON_ORACLE_EXCERPT = (
    "All three evaluator files contain labels for the same 1,074 items. Majority voting yielded 740 Unique and 334 Similar cases. "
    "Overall Fleiss' kappa was 0.1864; pairwise Cohen kappas were 0.1638, 0.1446, and 0.2671. Independence is not claimed."
)

reviewers = {
"Reviewer 1": [
 item("The generated and optimized tests are not executed; add execution success, coverage, mutation, fault-detection, and invalid-test measures, or narrow the conclusions.", COMMON_EXECUTION, "Narrowed the abstract, discussion, threats to validity, and conclusion; added a concrete execution-validation agenda.", "Abstract; Sections 6, 6.1, and 7", COMMON_EXECUTION_EXCERPT, "Partially addressed—claims narrowed; new execution experiments were not available"),
 item("The human reference uses one expert and lacks inter-rater and item-level evaluation.", COMMON_ORACLE, "Replaced the single-oracle analysis with three-evaluator consensus; added agreement and item-level metrics plus downloadable CSVs.", "Sections 4.3 and 5.4; Tables 12–14", COMMON_ORACLE_EXCERPT),
 item("Selecting the closest optimizer after examining results introduces selection bias.", "We agree. The revised manuscript reports every matched optimizer configuration and describes the old 96.52% figure only as a post-hoc count-level comparison.", "Removed best-only emphasis and added configuration-level matrices and metrics for 72 archived configurations.", "Section 5.4; revision_output/analysis/optimizer_item_level_metrics.csv", "The earlier 96.52% value compared only aggregate Unique/Similar counts after choosing the closest serial optimizer separately for each source. It is therefore a post-hoc count-level distributional similarity, not item-level accuracy and not a deployment selection rule."),
 item("Report repeated runs, seeds, decoding settings, versions, hardware/API details, token and concurrency limits, and statistical uncertainty.", COMMON_STOCHASTIC, "Added explicit reproducibility limitations and a preregistered repeated-run plan.", "Sections 6.1 and 7", COMMON_STOCHASTIC_EXCERPT, "Partially addressed—transparent limitation; no new repeated runs invented"),
 item("Explain merging across batches/parallel requests, conflicts, transitivity, representatives, and order effects.", "We now specify the archived decision rule and distinguish verified behaviour from unavailable production logic. Serial selection is order-dependent. Bulk logs describe single-call operations; over-context handling is unavailable. Parallel conflict, transitivity, scheduling, and merge policies could not be verified.", "Added Algorithm 1 and implementation-boundary text; identified unresolved merge semantics as a limitation.", "Section 3.3.7 and Algorithm 1; Sections 6 and 6.1", "Serial selection is order-sensitive. Cross-batch reconciliation for Bulk and conflict/transitivity handling for Parallel could not be verified from the supplied implementation evidence.", "Partially addressed—Serial verified; Bulk/Parallel production merge logic unavailable"),
 item("Clarify SCP and SRG with definitions, edge cases, and reproducible calculations.", "We derived the algebraic relationship SRG = SCP/(1−TRR), clarified that SRG is composite rather than independent, and added worked edge cases showing that high SRG can conceal scenario loss.", "Added a worked numerical paragraph immediately after the SCP equation and retained full configuration outputs.", "Section 5.3.1; Tables 8–11", "Because TRR = 1 − N_unique/N_generated and SCP = S_after/S_before, SRG is the composite SRG = SCP/(1 − TRR) when SCP is expressed as a fraction. It contains no information independent of SCP and TRR."),
 item("Reconcile model totals, scenarios, configurations, comparisons, and statistics with raw data and scripts.", "We audited the manuscript against the archived JSON and CSV data, corrected the GPT-OSS self-comparison arithmetic, and generated traceable item-level tables. We also reconciled the 74 non-empty archived records with the 72-configuration item-level set: record 0 is a 49-case partial suite and record 18 is an earlier auxiliary GPT-OSS/GPT-OSS Serial result whose systematically named counterpart is record 39.", "Corrected 137 = 80 + 57; documented the archive, main-table, and item-level scopes; added a numerical traceability audit and derived CSVs.", "Sections 3.3.7 and 5.4; Table 2; FINAL_NUMERICAL_TRACEABILITY_AUDIT.md; revision_output/analysis", "The archive contains 74 non-empty optimization records. Seventy-two form the item-level analysis set ... record 0 is an auxiliary 49-case partial suite ... record 18 is retained in the archive count but not counted as a separate benchmark configuration."),
 item("One robotic benchmark cannot support broad industrial generalization.", "We agree and explicitly restrict all conclusions to the archived robotic benchmark.", "Narrowed claims throughout and added benchmark/external-validity limitations.", "Abstract; Sections 4.1, 6.1, and 7", "The evidence comes from one industrially motivated robotic benchmark represented by requirements, UML/XML, and source code. Additional domains, programming languages, system sizes, architectures, and testing contexts are required before broader generalization."),
 item("Do not imply testing effectiveness or scalability when only pre-execution redundancy and representation are evaluated.", "We revised the claims to distinguish semantic redundancy, scenario representation, compactness, and traceability from executable correctness or industrial effectiveness.", "Rewrote the abstract, discussion, and conclusion using pre-execution language.", "Abstract; Sections 6 and 7", "They do not support claims about executable correctness, coverage, mutation score, fault detection, or general industrial effectiveness."),
 item("Shorten repetitive Serial/Bulk/Parallel descriptions and standardize terminology.", "We substantially compressed the manuscript and consolidated the strategy descriptions. Unique, Similar, redundant, and representative are now defined at the decision rule.", "Reduced both compiled manuscript variants from 49 to 31 pages and removed repeated methodological prose.", "Sections 1–7, especially Sections 3.3.7 and 5.2", "A case marked Similar is treated as semantically redundant under the recorded comparison decision; a retained Unique case is the representative for later comparisons."),
 item("Provide prompt templates and exact model settings; standardize model names/versions.", "The archived prompt semantics and observed model identifiers are now described, but complete production settings were not present and are identified as unavailable.", "Added prompt-field details, fixed model notation where data allowed, and documented missing snapshots/decoding parameters.", "Sections 3.2, 3.3.7, 4.3, and 6.1", "The archive preserves prompts, model identifiers, timestamps, and outputs, but temperatures, seeds, decoding parameters, hardware details, concurrency limits, and model snapshot hashes are not consistently recorded." , "Partially addressed—archive limitation disclosed"),
 item("The test-case schema lacks conventional preconditions, steps, data, expected results, and explicit oracles.", "We now characterize Title, Description, and Objective as an intermediate pre-execution schema and explicitly deny direct executability claims.", "Added a schema and optimization-validity limitation.", "Section 6.1", "Title, Description, and Objective form an intermediate pre-execution schema; preconditions, steps, data, expected results, and executable oracles may be embedded in prose rather than mandatory fields. Direct executability is not claimed."),
],
"Reviewer 2": [
 item("Reference 48 is an unverified preprint.", "We removed the preprint because it was not necessary to support the revised, narrower related-work claim.", "Removed the entry and its use from the literature review.", "Section 2 and References", "The staged design is relevant because unrestricted pairwise comparison grows quadratically."),
 item("References 59 and 63 lack publication metadata.", "The incomplete entries were not needed by the revised argument and were removed rather than completed speculatively.", "Removed the incomplete bibliography records.", "References", "Unsupported or incomplete entries were removed; the remaining cited entries compile without unresolved citations."),
 item("Specify the token-aware switching threshold and its consistency effects.", "The archive did not contain production configuration evidence for the stated threshold, fallback model, or switching scope. We therefore removed the unverified 4,000-token claim and state the reproducibility limitation.", "Replaced the threshold claim with an evidence-bounded description.", "Sections 3.2 and 4.2.1", "The revised study therefore does not claim a verified automatic threshold or fallback model; any deployment-specific model switching may introduce consistency effects that require separate evaluation.", "Partially addressed—unverifiable implementation claim removed"),
 item("Explain how textual inputs map to the binary semantic output.", "We now define the input fields, prompt priority, JSON Boolean output, and Serial decision procedure.", "Added a formal decision function and pseudocode.", "Section 3.3.7 and Algorithm 1", "The stored comparison prompt instructs the LLM to prioritize Description, then Objective, then Title; two cases are considered equivalent only when their validation intent and scenario are substantially the same."),
 item("The single benchmark limits generalizability.", "We agree and restrict the conclusions accordingly.", "Expanded the benchmark and external-validity threat.", "Sections 4.1, 6.1, and 7", "Additional domains, programming languages, system sizes, architectures, and testing contexts are required before broader generalization."),
 item("A single human expert introduces bias.", COMMON_ORACLE, "Used all three archived evaluator sets and quantified agreement.", "Sections 4.3 and 5.4; Table 12", COMMON_ORACLE_EXCERPT),
 item("Remove marketing-oriented phrasing in Section 5.2.", "We rewrote the opening of Section 5.2 so that it describes only the evaluated archived semantic decision process.", "Removed unsupported references to completeness, coverage breadth, expected results, preconditions, post-conditions, quality indexes, and execution adequacy.", "Section 5.2", "The evaluated Smart Selection configurations classify generated test cases as retained (Unique) or removed (Similar) using the semantic comparison procedure described in Section 3.3.7."),
 item("Aggregate alignment is insufficient; report item-level precision and recall.", "We agree. We report item-level precision, recall, F1, balanced accuracy, confusion counts, Cohen's kappa, false positives, and false negatives for every included configuration, and explicitly explain why the item-level set contains 72 of the 74 non-empty archived records.", "Added Tables 13–14, machine-readable outputs, and record-level inclusion/exclusion traceability.", "Section 5.4; FINAL_NUMERICAL_TRACEABILITY_AUDIT.md; revision_output/analysis", "The archive contains 74 non-empty optimization records. Seventy-two form the item-level analysis set."),
 item("The practical time analysis uses hypothetical intervals.", "We retain 78.9 hours only as a theoretical exhaustive-pair upper bound. We withdrew the prior 13-minute Bulk estimate because the archive has no elapsed-time field and its Bulk logs identify a single LLM call rather than 13 sequential batches.", "Rewrote the practical-time method and results and added a scope table.", "Sections 4.4 and 5.5; Table 17", "The previously stated 13-minute Bulk estimate is withdrawn: the archived Bulk logs contain no elapsed-time field and identify the evaluated operation as a single LLM call, not 13 sequential batches.", "Partially addressed—unsupported runtime estimate withdrawn; no measured timing study available"),
 item("Add a baseline against traditional heuristic optimization.", "We added a transparent, reproducible offline TF-IDF plus cosine-similarity baseline at three fixed thresholds. The archive does not establish prospective threshold selection, so we removed 'pre-specified' and do not claim that this single baseline is a comprehensive state-of-the-art comparison.", "Implemented and reported threshold-level item metrics and predictions; clarified the threshold-selection evidence boundary.", "Sections 4.3 and 5.4; Table 14; revision_output/analysis", "Three fixed analysis thresholds (0.30, 0.50, and 0.70) are reported. The supplied archive does not establish that they were selected before Human Oracle results were examined.", "Partially addressed"),
 item("Discuss the consequences of excluding execution validation and future plans.", COMMON_EXECUTION, "Added explicit implications and proposed execution, coverage, mutation, and traceability comparisons for retained versus removed tests.", "Sections 6, 6.1, and 7", COMMON_EXECUTION_EXCERPT),
],
"Reviewer 3": [
 item("What is new beyond LLM semantic deduplication? Compare with a simpler method.", "We sharpened the contribution as an integrated, traceable evaluation workflow rather than a new classifier and added the TF-IDF baseline.", "Rewrote the research gap/contributions and added the reproducible non-LLM baseline.", "Sections 1, 2, 4.3, and 5.4", "The study does not claim a new trained classifier, state-of-the-art superiority, or general effectiveness across software domains."),
 item("Table 2 has GPT-OSS-20B totals 137 versus 88 + 57 = 145.", "Thank you for identifying the inconsistency. The archived counts support U = 80 and S = 57, so T = 137.", "Corrected the table entry from 88 to 80.", "Table 2", "T = U + S."),
 item("Table 3's maximum and minimum comparison counts come from different source sets; why call their difference gain?", "We agree. We renamed the quantity cross-configuration comparison-count spread and removed the causal improvement and cost-reduction interpretation.", "Revised the equation label, table caption, column heading, and surrounding interpretation.", "Section 5.2.1 and Table 3", "Because these extrema may come from different optimizer configurations and source-suite conditions, the ratio is a descriptive heterogeneity measure; it is not a within-test-set causal gain, cost reduction, or effectiveness estimate."),
 item("SRG is composed only of SCP and TRR; what additional information does it provide?", "We agree that it contains no independent information. The revised text states this algebraically and treats SRG only as a compact trade-off summary.", "Added derivation and worked examples.", "Section 5.3.1", "It contains no information independent of SCP and TRR; its purpose is only to summarize their trade-off."),
 item("The closest optimizer cannot be selected without the oracle; show all optimizer results.", "We agree. The old closest-count calculation is now explicitly post hoc, and all 72 matched configurations are reported.", "Added full configuration metrics and decision CSVs.", "Section 5.4; revision_output/analysis", "The configuration-level confusion matrices and metrics are supplied in optimizer_item_level_metrics.csv; item decisions are supplied in optimizer_item_level_decisions.csv."),
 item("How often are distinct tests incorrectly removed?", "We define a false removal as a consensus-Unique item classified Similar and report the count for each configuration and baseline threshold.", "Added false-removal ranges and item-level outputs.", "Section 5.4; Tables 13–14", "Consensus-Unique cases incorrectly removed ranged from 0 to 65 for Serial, 0 to 57 for Bulk, and 10 to 41 for Parallel."),
 item("Strategies use different optimizer sets; compare common optimizers and report preservation for Bulk/Parallel.", "We agree that the archive is not a fully crossed strategy–optimizer experiment. We identified one limited same-source subset: the 169-case LLaMa3.2-3B suite evaluated with Gemini-2.5 Flash and Pro under Serial, Bulk, and Parallel. We report all six archived outcomes descriptively, while withholding a causal strategy ranking because they are single records and provider/run settings are incomplete.", "Added an archived-design table, a six-row common-optimizer table, configuration-level outputs, and an explicit confounding limitation.", "Sections 3.3.7, 5.4, and 6.1; Tables 2 and 16; revision_output/analysis", "These are single archived records, not repeated or provider-controlled trials, and therefore provide a descriptive common-optimizer check rather than a causal estimate of strategy effect.", "Partially addressed—limited common-optimizer evidence added; full controlled factorial comparison unavailable"),
 item("Order may change Serial results; explain duplicate merging and parallel conflicts.", "We confirm Serial order dependence. Bulk logs identify single-call operations, so the earlier cross-batch description was removed; behaviour beyond the archived input sizes remains unknown. Parallel conflict, transitivity, scheduling, and merge policies could not be verified.", "Added Algorithm 1, corrected the Bulk description, and expanded the limitations.", "Sections 3.3.7 and 6.1", "Because the representative set depends on earlier decisions, changing input order can change the result.", "Partially addressed—Serial verified; unavailable Bulk/Parallel implementation details disclosed"),
 item("Section 5.2 says experiments were repeated; report the number of runs.", "That statement was unsupported by the supplied archive and has been removed. The revised manuscript states that each reported configuration is represented by one archived output and therefore makes no stability or uncertainty claim.", "Replaced the repetition claim and added stochastic-validity limits plus a repeated-run future-work plan.", "Sections 5.2, 6.1, and 7", COMMON_STOCHASTIC_EXCERPT, "Partially addressed—contradiction removed; repeated-run robustness remains future work"),
 item("Listing 2 lacks inputs and expected results; clarify its purpose.", "We retain it only as an illustration of the intermediate generated schema and explicitly acknowledge that it is not an executable test specification.", "Added schema limitations and narrowed the listing's interpretation.", "Sections 4.2.2 and 6.1", "Title, Description, and Objective form an intermediate pre-execution schema ... Direct executability is not claimed."),
 item("Distinguish the 80-hour theoretical upper bound from the 13-minute computation estimate and human effort.", "We retain 78.9 hours only as a hypothetical exhaustive-pair upper bound. We withdrew the 13-minute computation estimate because the raw Bulk logs have no elapsed-time field and contradict the stated 13-batch execution pattern.", "Rewrote the practical-time method and results and added an evidence-scope table.", "Sections 4.4 and 5.5; Table 17", "The previously stated 13-minute Bulk estimate is withdrawn: the archived Bulk logs contain no elapsed-time field and identify the evaluated operation as a single LLM call, not 13 sequential batches.", "Partially addressed—unsupported estimate removed; no measured human or mode timing available"),
 item("Gemini Pro has lower Bulk TRR for only 3 of 7 source models, not most.", "Corrected exactly as noted and removed the associated qualitative ranking.", "Replaced 'most' with the exact three-of-seven comparison and interpreted it only as a recorded classification proportion.", "Section 5.2.2", "Gemini-2.5 Pro has lower TRR than Flash for three of the seven source suites in the archived Bulk comparison."),
 item("References 26 and 27 cite whole conferences; References 43–45 do not support the transformation claim.", "We removed these entries and rewrote the literature review so the claims rely only on aligned sources.", "Pruned unsupported whole-proceedings and misaligned citations.", "Section 2 and References", "These studies motivate traceable artifact transformation but do not by themselves validate generated tests."),
 item("Remove repeated strategy descriptions.", "We consolidated the definitions in the method and shortened repeated result narration.", "Compressed Sections 3–6 and standardized terminology.", "Sections 3.3.7 and 5.2", "Serial, Bulk, and Parallel are operational processing modes with different evidence gaps."),
 item("Figures 8–12 use text that is too small.", "We reviewed the final PDF page by page. The figures remain legible at the journal page size, while large matrices are retained in compact form because the source image assets did not support lossless label regeneration.", "Shortened the surrounding text and visually verified the final PDFs.", "Figures 8–12", "Complete item-level numerical outputs are additionally supplied as CSV files so interpretation does not depend on reading plotted labels alone.", "Partially addressed—visual QA completed; underlying raster assets limit relabelling"),
],
"Reviewer 4": [
 item("Shorten and reorganize the 49-page manuscript; remove repetition and move detail out of the main narrative.", "We substantially compressed and reorganized the paper while preserving the method, core results, and limitations.", "Reduced both compiled manuscript variants from 49 to 31 pages and moved item-level detail to reproducible CSV outputs.", "Entire manuscript; revision_output/analysis", "The complete implementation and item-level outputs are included in the revision analysis package."),
 item("Sharpen the objective and novelty with focused research questions.", "We now position the contribution as an integrated workflow and evaluation, not as a new semantic classifier, and state four research questions.", "Rewrote the abstract, introduction, and related work.", "Abstract; Sections 1 and 2", "The study asks four research questions. RQ1 examines how generation volume varies across artefact and model combinations ... RQ4 compares the LLM-mediated decisions with a reproducible TF-IDF and cosine-similarity baseline."),
 item("Characterize the dataset and explain development/validation/test separation.", "We characterized the benchmark using the archived generation records. The original XML, requirements, and source files are not included in the supplied revision package, so exact structural-element and requirement counts were removed rather than repeated without traceable source artefacts. The archive also does not establish a formal development-validation-test split.", "Reported traceable generation configuration counts, removed unverifiable input-structure counts, and disclosed the absent formal split.", "Sections 4.1, 5.1, and 6.1; FINAL_NUMERICAL_TRACEABILITY_AUDIT.md", "Because the original source artefacts are not included in the supplied revision package, exact requirement, sensor, branch, and structural-element counts are not asserted here.", "Partially addressed"),
 item("Repeat every stochastic configuration and report uncertainty.", COMMON_STOCHASTIC, "Added the missing-settings limitation and a specific preregistered repeated-run agenda.", "Sections 6.1 and 7", COMMON_STOCHASTIC_EXCERPT, "Partially addressed—no new stochastic reruns were available"),
 item("The Human Oracle lacks multiple independent reviewers and agreement.", COMMON_ORACLE, "Used the three archived evaluator files and reported agreement; independence, mutual label visibility, experience, blinding, and adjudication were not archived and are disclosed.", "Sections 4.3, 5.4, and 6.1", "The files document separate completion artefacts, but they do not record whether evaluators worked without interaction or could see one another's labels ... independence is therefore not claimed.", "Partially addressed"),
 item("The 96.52% count agreement is misleading; report item-level confusion metrics.", "We agree and have demoted the 96.52% value to explanatory context only.", "Added item-level confusion matrices, precision, recall, F1, balanced accuracy, kappa, and false-removal risk.", "Section 5.4; Tables 13–14; revision_output/analysis", "Identical class totals can conceal different item decisions. The new item-level analysis above is the primary evaluation; the count-level value is retained only to explain the earlier report."),
 item("Discuss semantic retrieval and reranking, including DOI 10.1371/journal.pone.0344683.", "We added a concise discussion of candidate retrieval followed by fine-grained reranking as a traceable alternative to unrestricted pairwise comparison and cited the specified peer-reviewed paper as an architectural analogy.", "Added the literature paragraph and complete PLOS ONE reference.", "Section 2; Reference 61", "A retrieval-reranking separation has been shown to improve efficiency and traceability in another high-stakes document-matching setting; we cite it as an architectural analogy rather than direct evidence for test optimization."),
 item("Redundancy reduction does not establish test quality; execute retained and removed tests.", COMMON_EXECUTION, "Narrowed all effectiveness claims and specified retained-versus-removed execution, coverage, mutation, fault, and traceability evaluation as future work.", "Sections 6, 6.1, and 7", COMMON_EXECUTION_EXCERPT, "Partially addressed—claims narrowed; execution study remains future work"),
 item("Compress large rotated tables and standardize figure presentation.", "We compressed the manuscript and retained only the decision-relevant aggregate tables in the main narrative; complete item-level results are delivered as CSVs. The remaining rotated tables are legacy aggregate matrices and were visually checked for clipping and readability.", "Reduced page count, shortened narration, pruned references, and verified every PDF page.", "Tables 2–15; revision_output/analysis", "Complete configuration-level confusion matrices and metrics are supplied in optimizer_item_level_metrics.csv; item decisions are supplied in optimizer_item_level_decisions.csv.", "Partially addressed—core aggregate tables retained for traceability"),
]}

if not REFERENCE.exists():
    raise FileNotFoundError(f"Reference response document not found: {REFERENCE}")


def clear_document_body(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def format_paragraph(paragraph, *, after=6, before=0, keep=False):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.keep_with_next = keep
    return paragraph


def add_labelled(document, label, text, *, green=False, keep=False, after=6):
    paragraph = format_paragraph(document.add_paragraph(), after=after, keep=keep)
    label_run = paragraph.add_run(label + " ")
    label_run.bold = True
    if green:
        label_run.font.color.rgb = RGBColor.from_string("008000")
    text_run = paragraph.add_run(text)
    if green:
        text_run.font.color.rgb = RGBColor.from_string("008000")
    return paragraph


with zipfile.ZipFile(REFERENCE) as package:
    separator_data = package.read("word/media/image1.png")

doc = Document(REFERENCE)
clear_document_body(doc)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.15

# Correspondence page: use the current manuscript identity and response text in
# the retained reference document's first-page pattern.
add_labelled(doc, "Original Manuscript ID:", "applsci-4556799", keep=True)
add_labelled(
    doc,
    "Original Article Title:",
    "STLC Manager: Benchmarking LLM-Driven Model-to-Test Transformation and Test Case Optimization",
    keep=True,
)
add_labelled(doc, "To:", "Editor, Applied Sciences", keep=True)
add_labelled(doc, "Re:", "Response to Reviewers", keep=True, after=12)
format_paragraph(doc.add_paragraph("Dear Editor,"), after=12)
format_paragraph(
    doc.add_paragraph(
        "We thank the reviewers for the detailed and constructive comments. We audited the supplied manuscript, raw JSON records, "
        "three Human Oracle directories, and archived optimization outputs before revising the paper. The revision uses no fabricated "
        "experiments. When requested evidence was absent, we narrowed the claims and marked the item as partially addressed. "
        "Missing evidence included repeated stochastic runs, execution outcomes, complete deployment configuration, evaluator biographies, and a balanced benchmark expansion."
    ),
    after=12,
)
format_paragraph(doc.add_paragraph("Best regards,"), after=0)
format_paragraph(doc.add_paragraph("The Authors"), after=0)
doc.add_page_break()

for reviewer_number, (_, entries) in enumerate(reviewers.items(), start=1):
    for comment_number, entry in enumerate(entries, start=1):
        prefix = f"Reviewer #{reviewer_number}, Comment #{comment_number}. Reviewer comment:"
        add_labelled(doc, prefix, entry["comment"], green=True, keep=True, after=8)
        add_labelled(doc, "Response:", entry["response"])
        add_labelled(doc, "Action taken:", entry["action"])
        add_labelled(doc, "Location:", entry["location"])
        add_labelled(doc, "Added or revised text:", f'“{entry["excerpt"]}”')
        add_labelled(doc, "Status:", entry["status"], after=7)

        is_last = reviewer_number == len(reviewers) and comment_number == len(entries)
        if not is_last:
            divider = doc.add_paragraph()
            divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
            divider.paragraph_format.space_before = Pt(2)
            divider.paragraph_format.space_after = Pt(7)
            divider.add_run().add_picture(BytesIO(separator_data), width=Inches(6.69), height=Inches(0.02))

for section in doc.sections:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for paragraph in section.header.paragraphs + section.footer.paragraphs:
        paragraph.clear()

doc.core_properties.title = "Response to Review STLC Manager"
doc.core_properties.subject = "Point-by-point responses to four reviewers"
doc.core_properties.author = "Authors"
doc.save(OUT)
print(f"Wrote {OUT} with {sum(map(len, reviewers.values()))} point-by-point responses")
